import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_proposal_docx(md_path, docx_path):
    doc = Document()

    # Set Margins to 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x24, 0x21) # dark charcoal/black
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        # Parse table lines
        header_line = tbl_lines[0].strip()
        data_lines = [l.strip() for l in tbl_lines[2:]] if len(tbl_lines) > 2 else []

        headers = [c.strip() for c in header_line.strip('|').split('|')]
        rows_data = []
        for dl in data_lines:
            if dl.startswith('|'):
                row = [c.strip() for c in dl.strip('|').split('|')]
                rows_data.append(row)

        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Header Row
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            shading = parse_xml(r'<w:shd {} w:fill="EAECEF"/>'.format(nsdecls('w')))
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)
            for p in hdr_cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10.5)
                    run.font.bold = True

        # Data Rows
        for r_idx, row in enumerate(rows_data):
            row_cells = table.rows[r_idx + 1].cells
            for c_idx, text in enumerate(row):
                if c_idx < len(row_cells):
                    row_cells[c_idx].text = text
                    for p in row_cells[c_idx].paragraphs:
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for run in p.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(10)

        # Borders
        tblPr = table._tbl.tblPr
        borders = parse_xml(r'<w:tblBorders {}><w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>'.format(nsdecls('w')))
        tblPr.append(borders)

        doc.add_paragraph() # Spacing after table

    i = 0
    while i < len(lines):
        line = lines[i]
        s_line = line.strip()

        # Handle Tables
        if s_line.startswith('|') and '|' in s_line[1:]:
            table_lines.append(line)
            in_table = True
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        if not s_line:
            i += 1
            continue

        # Page Breaks on major headings
        if s_line.startswith('# CHAPTER') or s_line in ['## DECLARATION', '## DEDICATION', '## ACKNOWLEDGEMENTS', '## TABLE OF CONTENTS', '## LIST OF TABLES', '## ABBREVIATIONS AND ACRONYMS', '## OPERATIONAL DEFINITION OF TERMS', '## ABSTRACT', '# REFERENCES']:
            doc.add_page_break()

        # Headings
        if s_line.startswith('# '):
            text = s_line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ('OPEN UNIVERSITY' in text or 'VIRTULAB' in text or 'REFERENCES' in text or 'CHAPTER' in text) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0C, 0x23, 0x40) # Navy

        elif s_line.startswith('## '):
            text = s_line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ('SCHOOL OF' in text or text in ['DECLARATION', 'DEDICATION', 'ACKNOWLEDGEMENTS', 'TABLE OF CONTENTS', 'LIST OF TABLES', 'ABBREVIATIONS AND ACRONYMS', 'OPERATIONAL DEFINITION OF TERMS', 'ABSTRACT']) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0C, 0x23, 0x40)

        elif s_line.startswith('### '):
            text = s_line[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if 'DEPARTMENT OF' in text else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True

        elif s_line.startswith('#### '):
            text = s_line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.italic = True

        elif s_line.startswith('---'):
            pass

        elif s_line.startswith('> '):
            text = s_line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.italic = True

        else:
            # Regular text paragraph
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)

            # Center title page details
            if s_line.startswith('**BY**') or s_line.startswith('**HARRISON') or s_line.startswith('**REG') or s_line.startswith('**SEPTEMBER') or 'Submitted in Partial Fulfilment' in s_line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Parse inline formatting (**bold**, *italic*)
            # Regex splitter
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', s_line)
            for part in parts:
                if not part:
                    continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.italic = True
                else:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)

        i += 1

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Document successfully created at {docx_path}")

if __name__ == '__main__':
    md_file = os.path.join(os.getcwd(), 'docs', 'PROJECT_PROPOSAL.md')
    docx_file = os.path.join(os.getcwd(), 'VirtuLab_Kenya_Capstone_Proposal_V8.docx')
    create_proposal_docx(md_file, docx_file)
