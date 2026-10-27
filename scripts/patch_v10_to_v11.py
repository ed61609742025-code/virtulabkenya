import os
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

v10_path = Path(r"c:\Users\USER\.gemini\antigravity\scratch\virtulabkenya\docs\capstone_proposals\VirtuLab_Kenya_Capstone_Proposal_V10_Resolved.docx")
v11_path = Path(r"c:\Users\USER\.gemini\antigravity\scratch\virtulabkenya\docs\capstone_proposals\VirtuLab_Kenya_Capstone_Proposal_V11_Resolved.docx")
downloads_v11 = Path(r"C:\Users\USER\Downloads\VirtuLab_Kenya_Capstone_Proposal_V11_Resolved.docx")

print(f"Loading {v10_path}...")
doc = docx.Document(str(v10_path))

# Helper to set cell styling
def style_header_cell(cell, text):
    cell.text = text
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79"/>')
    tcPr.append(shd)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9.5)

def style_body_cell(cell, text, bg="FFFFFF"):
    cell.text = text
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
    tcPr.append(shd)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            r.font.size = Pt(9.5)

# Text replacements across all paragraphs
replacements = [
    # 1. Word template residue
    ("Click (or Ctrl + Click) on any section or page number below to jump directly to it in this document. When opened in Microsoft Word, you can also right-click and choose 'Update Field' to synchronize entries.", ""),
    # 2. Title & Acronyms
    ("FOR KCSE AND CBE LEARNERS IN KENYAN SECONDARY SCHOOLS", "FOR KCSE AND CBC/CBE LEARNERS IN KENYAN SECONDARY SCHOOLS"),
    ("ADDE (Analyse, Design, Develop, and Evaluate)", "ADDIE (Analysis, Design, Development, Implementation, and Evaluation)"),
    # 3. Node.js version
    ("Node.js version 18", "Node.js 20 LTS"),
    ("Node.js 18", "Node.js 20 LTS"),
    # 4. Third-person voice
    ("I will also work with County Education offices and CEMASTEA", "The researcher will also collaborate with County Education offices and CEMASTEA"),
    # 5. Governance accuracy
    ("To County Education Boards: Allocate targeted capitation grants toward procuring low-cost Android tablets", 
     "To the Ministry of Education (State Department for Basic Education) and NG-CDF Committees: Earmark targeted ICT infrastructure and capitation grants toward procuring low-cost Android tablet kits"),
    # 6. Retrospective claim reframing
    ("Drawing on pilot findings, the study presents three actionable policy recommendations",
     "Contingent upon validating the target learning gain (g >= 0.40) and usability (SUS >= 78) benchmarks during the field pilot, the study advances three actionable policy recommendations"),
    ("VirtuLab Kenya demonstrates that thoughtful instructional design",
     "VirtuLab Kenya is engineered to demonstrate that thoughtful instructional design"),
    # 7. Undergraduate titration purging
    ("precipitation, and complexometric titrations", "back-titrations, and water of crystallisation determinations"),
    ("complexometric and precipitation titrations", "back-titrations, double-indicator, and redox titrations"),
    # 8. Past papers bank
    ("1989 to 2013", "2010 to 2025"),
    ("1989–2016", "2010–2025")
]

replaced_count = 0
for p in doc.paragraphs:
    for old_text, new_text in replacements:
        if old_text in p.text:
            p.text = p.text.replace(old_text, new_text)
            replaced_count += 1

print(f"Applied {replaced_count} inline text patches across paragraphs.")

# Text replacements across tables
table_replaced = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for old_text, new_text in replacements:
                    if old_text in p.text:
                        p.text = p.text.replace(old_text, new_text)
                        table_replaced += 1

print(f"Applied {table_replaced} text patches inside tables.")

# Save V11
doc.save(str(v11_path))
print(f"Saved primary V11 to: {v11_path}")

try:
    shutil.copyfile(str(v11_path), str(downloads_v11))
    print(f"Copied V11 directly to Downloads: {downloads_v11}")
except Exception as e:
    print(f"Could not copy to Downloads: {e}")

print("V11 RESOLVED DOCX GENERATED SUCCESSFULLY!")
