import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

brain_path = Path(r"C:\Users\USER\.gemini\antigravity\brain\1347d73b-10ce-4cb1-92ea-39097135637a\VirtuLab_Kenya_CEMASTEA_Collaboration_Letter.md")
docx_path = Path(r"c:\Users\USER\.gemini\antigravity\scratch\virtulabkenya\docs\VirtuLab_Kenya_CEMASTEA_Collaboration_Letter.docx")
downloads_path = Path(r"C:\Users\USER\Downloads\VirtuLab_Kenya_CEMASTEA_Collaboration_Letter.docx")

doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(4)

with open(brain_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_code = False
code_lines = []

for line in lines:
    raw = line.rstrip('\r\n')
    stripped = raw.strip()
    
    if stripped.startswith('```'):
        if in_code:
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            tcPr = cell._element.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F5F7"/>'))
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="100" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/><w:left w:w="150" w:type="dxa"/><w:right w:w="150" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run("\n".join(code_lines))
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            code_lines = []
            in_code = False
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        else:
            in_code = True
        continue
        
    if in_code:
        code_lines.append(raw)
        continue
        
    if stripped.startswith('# '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(stripped[2:])
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    elif stripped.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(stripped[4:])
        r.font.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    elif stripped.startswith('#### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(stripped[5:])
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x2E, 0x5B, 0x88)
    elif stripped.startswith(('1. ', '2. ', '3. ', '4. ', '5. ')):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(2)
        parts = stripped[3:].split('**')
        for idx, pt in enumerate(parts):
            r = p.add_run(pt)
            if idx % 2 == 1:
                r.font.bold = True
    elif stripped.startswith('* ') or stripped.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        parts = stripped[2:].split('**')
        for idx, pt in enumerate(parts):
            r = p.add_run(pt)
            if idx % 2 == 1:
                r.font.bold = True
    elif stripped == '---':
        p = doc.add_paragraph()
        r = p.add_run("―" * 45)
        r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    elif stripped:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        parts = stripped.split('**')
        for idx, pt in enumerate(parts):
            r = p.add_run(pt)
            if idx % 2 == 1:
                r.font.bold = True

doc.save(str(docx_path))
doc.save(str(downloads_path))
print("SAVED COLLABORATION LETTER DOCX TO BOTH DOCS AND DOWNLOADS!")
